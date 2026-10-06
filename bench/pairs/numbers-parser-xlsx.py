"""numbers-parser-xlsx.py <in.numbers> <out.xlsx>: the Excel reference for
the numbers-csv pair. Reads every table with numbers-parser and writes each
as a worksheet with openpyxl in write-only (streaming) mode, as `sublime`
writes a worksheet per table."""
import sys

from numbers_parser import Document
from openpyxl import Workbook

document = Document(sys.argv[1])
workbook = Workbook(write_only=True)
for sheet in document.sheets:
    for table in sheet.tables:
        name = sheet.name if len(sheet.tables) == 1 else f"{sheet.name} - {table.name}"
        worksheet = workbook.create_sheet(title=name[:31].replace("/", "_"))
        for row in table.rows(values_only=True):
            worksheet.append(list(row))
workbook.save(sys.argv[2])
